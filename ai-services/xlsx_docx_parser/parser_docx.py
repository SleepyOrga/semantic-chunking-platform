import logging
import time
from pathlib import Path
import os
import argparse
import tempfile
import urllib.request
import boto3
import json
import pika
from typing import Dict, Any, Optional

from docx import Document
from docx.document import Document as DocumentType

from docling.document_converter import DocumentConverter, WordFormatOption
from docling.datamodel.base_models import InputFormat
from docling.pipeline.simple_pipeline import SimplePipeline
from docling.backend.msword_backend import MsWordDocumentBackend

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

class DocxParserService:
    def __init__(self, rabbitmq_url: str):
        self.rabbitmq_url = rabbitmq_url
        self.connection = None
        self.channel = None
        self.s3_client = boto3.client('s3')

    def connect(self):
        """Establish connection to RabbitMQ."""
        try:
            self.connection = pika.BlockingConnection(
                pika.URLParameters(self.rabbitmq_url)
            )
            self.channel = self.connection.channel()
            
            # Declare queues
            self.channel.queue_declare(queue='docx-parser-queue', durable=True)
            self.channel.queue_declare(queue='chunking-queue', durable=True)
            
            logger.info("✅ Connected to RabbitMQ and listening on docx-parser-queue")
            
        except Exception as e:
            logger.error(f"Failed to connect to RabbitMQ: {e}")
            raise

def extract_images_from_docx(docx_path, images_dir, docx_stem):
    image_mapping = {}
    image_count = 0
    try:
        doc = Document(docx_path)
        for rel_id, rel in doc.part.rels.items():
            if "image" in rel.target_ref:
                try:
                    image_data = rel.target_part.blob
                    content_type = rel.target_part.content_type
                    ext = ".png"
                    if "jpeg" in content_type or "jpg" in content_type:
                        ext = ".jpg"
                    elif "gif" in content_type:
                        ext = ".gif"
                    elif "bmp" in content_type:
                        ext = ".bmp"

                    image_filename = f"{docx_stem}_image_{image_count + 1}{ext}"
                    image_path = os.path.join(images_dir, image_filename)

                    with open(image_path, "wb") as img_file:
                        img_file.write(image_data)

                    image_mapping[rel_id] = image_filename
                    image_count += 1
                    logging.info(f"🖼️ Extracted image: {image_filename}")
                except Exception as e:
                    logging.warning(f"Failed to extract image: {e}")
    except Exception as e:
        logging.warning(f"python-docx error: {e}")
    return image_mapping

#Upload images to S3 and return mapping of local filenames to S3 URLs
def upload_images_to_s3_and_get_urls(images_dir, s3_bucket, s3_prefix, image_mapping):
    """Upload images to S3 and return mapping of local filenames to S3 URLs - uses existing upload_to_s3"""
    s3_urls = {}
    if not os.path.isdir(images_dir):
        return s3_urls
    
    for rel_id, image_filename in image_mapping.items():
        local_image_path = os.path.join(images_dir, image_filename)
        if os.path.isfile(local_image_path):
            try:
                # Use existing upload_to_s3 function
                s3_image_key = f"{s3_prefix}images/{image_filename}"
                upload_to_s3(local_image_path, s3_bucket, s3_image_key)
                
                # Generate S3 URL
                s3_url = f"https://{s3_bucket}.s3.amazonaws.com/{s3_image_key}"
                s3_urls[image_filename] = s3_url
                
            except Exception as e:
                logging.error(f"Failed to upload image {image_filename}: {e}")
    
    return s3_urls
# Embedd S3 URLs in markdown image tag 
def process_markdown_with_s3_images(markdown_content, image_mapping, s3_image_urls):
    """Process markdown content with S3 URLs - leverages existing process_markdown_with_images logic"""
    if not s3_image_urls:
        # Fall back to existing function
        return process_markdown_with_images(markdown_content, image_mapping)

    # Use same logic as existing function but with S3 URLs
    modified_content = markdown_content
    
    # Replace placeholders with S3 URLs instead of local paths
    for rel_id, image_filename in image_mapping.items():
        if image_filename in s3_image_urls:
            s3_url = s3_image_urls[image_filename]
            image_ref = f"![{image_filename}]({s3_url})"
            modified_content = modified_content.replace("<!-- image -->", image_ref, 1)
    
    # Remove remaining placeholders (same as existing function)
    modified_content = modified_content.replace("<!-- image -->", "")
    
    # Add remaining images (same logic as existing function but with S3 URLs)
    remaining_images = [img for img in s3_image_urls.keys() if f"![{img}]" not in modified_content]
    if remaining_images:
        modified_content += "\n\n## Additional Extracted Images\n\n"
        for img_filename in remaining_images:
            s3_url = s3_image_urls[img_filename]
            modified_content += f"![{img_filename}]({s3_url})\n\n"

    return modified_content

def download_file_from_url(url: str) -> str:
    """
    Download a file from a URL to a temporary location.
    
    Args:
        url: The URL of the file to download
        
    Returns:
        str: Path to the downloaded temporary file
    """
    temp_fd, temp_path = tempfile.mkstemp(suffix=os.path.splitext(url)[1])
    os.close(temp_fd)  # Close the file descriptor
    logger.info(f"Downloading file from URL: {url}")
    try:
        urllib.request.urlretrieve(url, temp_path)
        logger.info(f"Successfully downloaded file to temporary location: {temp_path}")
        return temp_path
    except Exception as e:
        logger.error(f"Failed to download file from {url}: {e}")
        if os.path.exists(temp_path):
            os.unlink(temp_path)
        raise

def process_markdown_with_images(markdown_content, image_mapping):
    if not image_mapping:
        return markdown_content

    modified_content = markdown_content
    image_list = list(image_mapping.values())

    for i, image_filename in enumerate(image_list):
        image_ref = f"![{image_filename}](images/{image_filename})"
        modified_content = modified_content.replace("<!-- image -->", image_ref, 1)

    modified_content = modified_content.replace("<!-- image -->", "")

    remaining_images = image_list[modified_content.count("!["):]
    if remaining_images:
        modified_content += "\n\n## Additional Extracted Images\n\n"
        for img_file in remaining_images:
            modified_content += f"![{img_file}](images/{img_file})\n\n"

    return modified_content


def upload_to_s3(local_path, bucket, s3_key):
    s3 = boto3.client('s3')
    s3.upload_file(str(local_path), bucket, s3_key)
    logging.info(f"Uploaded {local_path} to s3://{bucket}/{s3_key}")


def extract_docx_to_markdown(input_path, output_dir, extract_images=True, s3_bucket=None, s3_prefix=None):
    start = time.time()
    os.makedirs(output_dir, exist_ok=True)
    if extract_images:
        images_dir = os.path.join(output_dir, "images")
        os.makedirs(images_dir, exist_ok=True)

    converter = DocumentConverter(
        format_options={
            InputFormat.DOCX: WordFormatOption(
                pipeline_cls=SimplePipeline,
                backend=MsWordDocumentBackend
            )
        }
    )

    input_path = Path(input_path)
    if input_path.is_file() and input_path.suffix.lower() == ".docx":
        files_to_process = [input_path]
    elif input_path.is_dir():
        files_to_process = list(input_path.glob("*.docx"))
    else:
        logging.error(f"Invalid input path: {input_path}")
        return

    logging.info(f"📄 Found {len(files_to_process)} DOCX files")

    for docx_file in files_to_process:
        try:
            logging.info(f"🚀 Processing {docx_file.name}")
            result = converter.convert(docx_file)

            output_name = docx_file.stem + ".md"
            output_path = os.path.join(output_dir, output_name)

            markdown_content = result.document.export_to_markdown()

            image_mapping = {}

            if extract_images:
                logging.info(f"Extracting images from {docx_file.name}...")
                image_mapping = extract_images_from_docx(docx_file, images_dir, docx_file.stem)                
                # If S3 bucket is provided, upload images and get S3 URLs
                if s3_bucket and s3_prefix and image_mapping:
                    logging.info(f"📤 Uploading {len(image_mapping)} images to S3...")
                    s3_image_urls = upload_images_to_s3_and_get_urls(
                        images_dir, s3_bucket, s3_prefix, image_mapping
                    )
                    # Use S3 URLs in markdown
                    final_markdown = process_markdown_with_s3_images(
                        markdown_content, image_mapping, s3_image_urls
                    )
                else:
                    # Use local image references
                    final_markdown = process_markdown_with_images(markdown_content, image_mapping)
            else:
                final_markdown = markdown_content

            with open(output_path, "w", encoding="utf-8") as f:
                f.write(final_markdown)
            logging.info(f"Saved markdown: {output_path}")

            # --- S3 upload logic ---
            if s3_bucket and s3_prefix:
                # Upload final embedded markdown
                md_s3_key = f"{s3_prefix}{Path(output_path).name}"
                upload_to_s3(output_path, s3_bucket, md_s3_key)
                # Print JSON so NestJS can capture it
                print(json.dumps({"md_s3_key": md_s3_key}))
        except Exception as e:
            logging.error(f"Error processing {docx_file.name}: {e}")

    logging.info(f"Completed in {time.time() - start:.2f} seconds")
    return md_s3_key

def download_file_from_url(url):
    # Tạo file tạm
    temp_fd, temp_path = tempfile.mkstemp(suffix=".docx")
    os.close(temp_fd)  # Đóng file descriptor
    logging.info(f"Downloading from URL: {url}")
    urllib.request.urlretrieve(url, temp_path)
    return temp_path

class DocxParserService:
    def __init__(self, rabbitmq_url: str):
        self.rabbitmq_url = rabbitmq_url
        self.connection = None
        self.channel = None
        self.s3_client = boto3.client('s3')

    def connect(self):
        """Establish connection to RabbitMQ."""
        try:
            self.connection = pika.BlockingConnection(
                pika.URLParameters(self.rabbitmq_url)
            )
            self.channel = self.connection.channel()
            
            # Declare queues
            self.channel.queue_declare(queue='docx-parser-queue', durable=True)
            self.channel.queue_declare(queue='chunking-queue', durable=True)
            
            logger.info("✅ Connected to RabbitMQ and listening on docx-parser-queue")
            
        except Exception as e:
            logger.error(f"Failed to connect to RabbitMQ: {e}")
            raise

    def process_message(self, ch, method, properties, body):
        s3 = boto3.client('s3')
        s3.upload_file(str(local_path), bucket, s3_key)
        logging.info(f"Uploaded {local_path} to s3://{bucket}/{s3_key}")

    def start_consuming(self):
        self.channel.basic_consume(queue='docx-parser-queue',
                                  on_message_callback=self.process_message,
                                  no_ack=True)

    def process_message(self, ch, method, properties, body):
        """Process incoming message from the queue."""
        try:
            payload = json.loads(body)
            logger.info(f"📩 Received message: {payload}")
            
            # Extract necessary information from payload
            input_path = payload.get('s3Key')
            filename = payload.get('filename', 'unknown')
            document_id = payload.get('documentId')
            s3_bucket = payload.get('s3Bucket') or os.getenv('S3_BUCKET_NAME', 'semantic-chunking-bucket')
            
            # Generate output path
            output_s3_prefix = f"parsed/{document_id}/"
            
            logger.info(f"🚀 Processing DOCX file: {filename}")
            
            # Call the parser function
            result = extract_docx_to_markdown(
                input_path=input_path,
                output_dir=output_s3_prefix,
                extract_images=True,
                s3_bucket=s3_bucket,
                s3_prefix=output_s3_prefix
            )
            
            # Send result to chunking queue
            chunking_payload = {
                's3Bucket': s3_bucket,
                's3Key': result,
                'documentId': document_id,
                'fileType': 'docx'
            }
            
            self.channel.basic_publish(
                exchange='',
                routing_key='chunking-queue',
                body=json.dumps(chunking_payload),
                properties=pika.BasicProperties(
                    delivery_mode=2,  # Make message persistent
                )
            )
            
            logger.info(f"✅ Successfully processed {filename}")
            ch.basic_ack(delivery_tag=method.delivery_tag)
            
        except Exception as e:
            logger.error(f"❌ Error processing message: {e}", exc_info=True)
            ch.basic_nack(delivery_tag=method.delivery_tag, requeue=False)

    def start_consuming(self):
        """Start consuming messages from the queue."""
        try:
            self.channel.basic_qos(prefetch_count=1)
            self.channel.basic_consume(
                queue='docx-parser-queue',
                on_message_callback=self.process_message,
                auto_ack=False
            )
            
            logger.info("🔄 Waiting for messages in docx-parser-queue...")
            self.channel.start_consuming()
            
        except KeyboardInterrupt:
            logger.info("\n👋 Shutting down DOCX parser...")
            if self.connection:
                self.connection.close()


def main():
    parser = argparse.ArgumentParser(description='Run DOCX parser worker')
    parser.add_argument('--rabbitmq', type=str, 
                       default='amqp://admin:admin@52.65.216.159:5672',
                       help='RabbitMQ connection URL')
    
    args = parser.parse_args()
    
    # Create and start worker
    worker = DocxParserService(args.rabbitmq)
    worker.connect()
    worker.start_consuming()

if __name__ == "__main__":
    main()